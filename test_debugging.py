"""Offline regression tests for API failures, batch isolation, and Drive/Excel I/O."""

import datetime
import io
import os
import unittest
import xml.etree.ElementTree as ET
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

import main


API = "http://schemas.nav.gov.hu/OSA/3.0/api"
COMMON = "http://schemas.nav.gov.hu/NTCA/1.0/common"


def company(code="TEST"):
    return dict(company_code=code, nav_login="test-user", nav_password="test-password",
                nav_tax_number="12345678", nav_signature_key="test-key",
                nav_base_url="https://example.invalid/invoiceService/v3",
                target_folder_id="test-folder", active=True)


def response(page, pages, number=None):
    digest = f"<invoiceDigest><invoiceNumber>{number}</invoiceNumber></invoiceDigest>" if number else ""
    return Mock(status_code=200, text=(
        f'<QueryInvoiceDigestResponse xmlns="{API}" xmlns:common="{COMMON}">'
        '<common:result><common:funcCode>OK</common:funcCode></common:result>'
        f'<currentPage>{page}</currentPage><availablePage>{pages}</availablePage>'
        f'{digest}</QueryInvoiceDigestResponse>'
    ))


class PaginationAndFailureTest(unittest.TestCase):
    def test_fetches_every_page_in_order(self):
        with patch.object(main.requests, "post", side_effect=[response(1, 3, "A"), response(2, 3, "B"), response(3, 3, "C")]) as post:
            rows, request, reply = main.fetch_all_invoices(company(), "2026-09-01", "2026-09-02")
        self.assertEqual(rows["invoiceNumber"].tolist(), ["A", "B", "C"])
        self.assertEqual([ET.fromstring(c.kwargs["data"]).findtext(f"{{{API}}}page") for c in post.call_args_list], ["1", "2", "3"])
        self.assertIn("<page>3</page>", request)
        self.assertIn("<invoiceNumber>C</invoiceNumber>", reply)

    def test_empty_result_stops_after_one_request(self):
        with patch.object(main.requests, "post", return_value=response(1, 0)) as post:
            rows, _, _ = main.fetch_all_invoices(company(), "2026-09-01", "2026-09-02")
        self.assertTrue(rows.empty)
        post.assert_called_once()

    def test_http_and_malformed_xml_failures_retain_diagnostics(self):
        for status, body in [(503, "unavailable"), (200, "<broken")]:
            with self.subTest(status=status), patch.object(main.requests, "post", return_value=Mock(status_code=status, text=body)):
                with self.assertRaises(RuntimeError) as caught:
                    main.fetch_all_invoices(company(), "2026-09-01", "2026-09-02")
            self.assertIn("QueryInvoiceDigestRequest", caught.exception.args[1])
            self.assertEqual(caught.exception.args[2], body)


class ConfigurationAndBatchTest(unittest.TestCase):
    def test_invalid_company_schemas_are_rejected(self):
        valid = pd.DataFrame([company()])
        for frame, message in [(valid.drop(columns="nav_password"), "missing columns"),
                               (valid.iloc[:0], "no rows"),
                               (pd.concat([valid, valid]), "unique"),
                               (valid.assign(active="yes"), "TRUE/FALSE")]:
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                main.validate_company_schema(frame)

    def test_loading_filters_inactive_companies(self):
        frame = pd.DataFrame([company("ON"), dict(company("OFF"), active=False)])
        stream = io.BytesIO()
        frame.to_excel(stream, sheet_name="companies", index=False)
        stream.seek(0)
        drive = Mock()
        drive.download_as_excel_stream.return_value = stream
        with patch.dict(os.environ, {"COMPANY_CONFIG_FILE_ID": "config"}), patch.object(main, "DriveClient", return_value=drive):
            loaded = main.load_companies_from_drive()
        self.assertEqual(loaded["company_code"].tolist(), ["ON"])
        drive.download_as_excel_stream.assert_called_once_with("config")

    def test_failure_isolated_and_diagnostics_logged_while_next_company_exports(self):
        good = pd.DataFrame({"invoiceNumber": ["TEST-1"], "invoiceNetAmount": ["100"],
                             "invoiceVatAmount": ["27"], "invoiceIssueDate": ["invalid"]})
        fixed_date = Mock(wraps=datetime.date)
        fixed_date.today.return_value = datetime.date(2026, 9, 15)
        with patch.dict(os.environ, {"COMPANY_CONFIG_FILE_ID": "config", "SUMMARY_LOG_FOLDER_ID": "summary"}), \
             patch.object(main, "load_companies_from_drive", return_value=pd.DataFrame([company("BAD"), company("GOOD")])), \
             patch.object(main, "fetch_all_invoices", side_effect=[RuntimeError("failure", "request-xml", "response-xml"), (good, "", "")]) as fetch, \
             patch.object(main, "upsert_company_excel") as upsert, \
             patch.object(main, "upload_summary_log") as summary, \
             patch.object(main.datetime, "date", fixed_date):
            result, status = main.weekly_invoice_export(None)
        self.assertEqual((result, status), ({"status": "ok", "companies": 2}, 200))
        self.assertEqual(fetch.call_args.args[1:], ("2026-09-07", "2026-09-13"))
        upsert.assert_called_once()
        exported = upsert.call_args.args[0]
        self.assertEqual(list(exported.columns), main.OUTPUT_COLUMNS)
        self.assertEqual(exported.loc[0, "invoiceGrossAmount"], 127)
        self.assertTrue(pd.isna(exported.loc[0, "invoiceIssueDate"]))
        log = summary.call_args.args[0]
        self.assertEqual(log["status"].tolist(), ["FAILED", "SUCCESS"])
        self.assertEqual(log.loc[0, "request_xml"], "request-xml")
        self.assertEqual(log.loc[0, "nav_error_response"], "response-xml")
        self.assertEqual(summary.call_args.args[1], "summary_2026-09-07_2026-09-13.xlsx")


class DriveAndExcelTest(unittest.TestCase):
    def make_drive(self):
        drive = main.DriveClient.__new__(main.DriveClient)
        drive.service = Mock()
        return drive

    def test_metadata_supports_shared_drive_and_permission_fields(self):
        drive = self.make_drive()
        drive.get_metadata("folder", fields="driveId,capabilities(canAddChildren)")
        drive.service.files().get.assert_called_once_with(fileId="folder", fields="driveId,capabilities(canAddChildren)", supportsAllDrives=True)

    def test_download_supports_google_sheets_and_binary_excel(self):
        for mime, export in [("application/vnd.google-apps.spreadsheet", True),
                             ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", False)]:
            with self.subTest(mime=mime):
                drive = self.make_drive()
                drive.service.files().get().execute.return_value = {"mimeType": mime}
                def downloader(stream, request):
                    stream.write(b"test content")
                    result = Mock()
                    result.next_chunk.return_value = (None, True)
                    return result
                with patch.object(main, "MediaIoBaseDownload", side_effect=downloader):
                    stream = drive.download_as_excel_stream("file")
                self.assertEqual(stream.read(), b"test content")
                if export:
                    drive.service.files().export.assert_called_once_with(fileId="file", mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    drive.service.files().get_media.assert_not_called()
                else:
                    drive.service.files().get_media.assert_called_once_with(fileId="file", supportsAllDrives=True)
                    drive.service.files().export.assert_not_called()

    def test_folder_search_includes_shared_drives(self):
        drive = self.make_drive()
        drive.service.files().list().execute.return_value = {"files": []}
        self.assertIsNone(drive.find_file_in_folder("test.xlsx", "folder"))
        kwargs = drive.service.files().list.call_args.kwargs
        self.assertTrue(kwargs["supportsAllDrives"])
        self.assertTrue(kwargs["includeItemsFromAllDrives"])

    def test_excel_cells_have_date_numeric_formats_and_capped_widths(self):
        frame = pd.DataFrame({"invoiceNumber": ["X" * 100], "invoiceIssueDate": [pd.Timestamp("2026-09-01")],
                              "invoiceNetAmount": [100.25], "invoiceGrossAmount": [127.32]}).reindex(columns=main.OUTPUT_COLUMNS)
        stream = io.BytesIO()
        main.write_excel_with_autowidth(frame, stream, max_width=30)
        book = load_workbook(stream)
        sheet = book.active
        for name in main.DATE_COLUMNS + main.NUMERIC_COLUMNS:
            index = main.OUTPUT_COLUMNS.index(name) + 1
            self.assertEqual(sheet.cell(2, index).number_format, "yyyy-mm-dd" if name in main.DATE_COLUMNS else "#,##0")
        self.assertEqual(sheet["H2"].value, 100.25)
        self.assertLessEqual(sheet.column_dimensions["B"].width, 30)

    def test_gross_preserves_zero_negative_and_invoice_currency(self):
        frame = pd.DataFrame({"invoiceNetAmount": [0, -100, 100.25], "invoiceVatAmount": [0, -27, 0],
                              "invoiceNetAmountHUF": [999, 999, 999], "invoiceVatAmountHUF": [999, 999, 999]})
        self.assertEqual(main.add_calculated_amounts(frame)["invoiceGrossAmount"].tolist(), [0, -127, 100.25])

    def test_excel_handles_missing_dates_strings_and_empty_results(self):
        frame = pd.DataFrame({
            "invoiceNumber": ["SYNTHETIC", None],
            "invoiceIssueDate": pd.to_datetime(["2026-09-01", None]),
        }).reindex(columns=main.OUTPUT_COLUMNS)
        for data in (frame, frame.iloc[:0]):
            with self.subTest(rows=len(data)):
                stream = io.BytesIO()
                main.write_excel_with_autowidth(data, stream)
                sheet = load_workbook(stream).active
                self.assertEqual([cell.value for cell in sheet[1]], main.OUTPUT_COLUMNS)
                self.assertGreaterEqual(sheet.column_dimensions["A"].width, len("invoiceIssueDate"))
                if len(data):
                    self.assertIsNone(sheet["A3"].value)
                    self.assertIsNone(sheet["B3"].value)


if __name__ == "__main__":
    unittest.main()
